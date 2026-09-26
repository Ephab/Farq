"""Build a read-only Blackboard demo snapshot from approved local lectures.

The script reads only five explicitly mapped course folders and uploads text to
Farq's authenticated internal import endpoint. Raw files never leave the host.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from pypdf import PdfReader
from pptx import Presentation


MAX_TEXT = 120_000
COURSES = (
    {"external_id": "ARTI-404", "code": "ARTI 404", "title": "Machine Learning", "folder": ("6th Semester", "ML", "Slides"), "recursive": True, "description": "Core machine-learning workflows, supervised models, ensembles, clustering, and density estimation."},
    {"external_id": "CYS-401", "code": "CYS 401", "title": "Cybersecurity", "folder": ("6th Semester", "CYS", "Slides"), "recursive": True, "description": "Authentication, access control, intrusion detection, secure design, and security management."},
    {"external_id": "ARTI-310", "code": "ARTI 310", "title": "Robotics", "folder": ("6th Semester", "Robotics", "Theory", "PDF_Chapters"), "recursive": True, "description": "Robot sensing, locomotion, control, representation, navigation, and multi-robot learning."},
    {"external_id": "CSC-311", "code": "CSC 311", "title": "Algorithms", "folder": ("6th Semester", "Algorithms"), "recursive": False, "description": "Algorithm analysis, divide and conquer, dynamic programming, greedy methods, trees, and graphs."},
    {"external_id": "ARTI-309", "code": "ARTI 309", "title": "Mathematics for AI", "folder": ("6th Semester", "Math4AI", "Lectures"), "recursive": True, "description": "Mathematical foundations used across machine learning and artificial intelligence."},
)
BLOCKED_NAME = re.compile(r"\b(solution|solved|answer|quiz|exam|final20\d\d|project|report)\b", re.IGNORECASE)


def read_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)[:MAX_TEXT]


def extract_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text.strip())
        if parts:
            slides.append(f"Slide {index}\n" + "\n".join(parts))
        if sum(len(item) for item in slides) >= MAX_TEXT:
            break
    return "\n\n".join(slides)[:MAX_TEXT]


def eligible_files(folder: Path, recursive: bool) -> list[Path]:
    candidates = folder.rglob("*") if recursive else folder.glob("*")
    return sorted(
        path for path in candidates
        if path.is_file()
        and path.suffix.lower() in {".pdf", ".pptx"}
        and not BLOCKED_NAME.search(path.stem)
    )


def synthetic_items(course: dict) -> list[dict]:
    code = course["code"]
    slug = course["external_id"].lower()
    return [
        {
            "external_id": f"{slug}-syllabus",
            "content_type": "syllabus",
            "title": f"{code} course syllabus (demo)",
            "body_text": f"Demo syllabus for {course['title']}. Weekly lectures, guided practice, one applied assessment, and a final review. This record is synthetic and exists only for the Farq hackathon demonstration.",
            "origin": "synthetic",
            "source_ref": f"bb://{course['external_id']}/syllabus",
            "posted_at": "2026-09-01T08:00:00Z",
            "modified_at": "2026-09-01T08:00:00Z",
        },
        {
            "external_id": f"{slug}-announcement-1",
            "content_type": "announcement",
            "title": "New lecture material available (demo)",
            "body_text": f"The latest {course['title']} lecture material is now available. Review it before the next class. This announcement is synthetic demo data.",
            "origin": "synthetic",
            "source_ref": f"bb://{course['external_id']}/announcements/1",
            "posted_at": "2026-09-24T09:00:00Z",
            "modified_at": "2026-09-24T09:00:00Z",
        },
        {
            "external_id": f"{slug}-assignment-1",
            "content_type": "assignment",
            "title": f"{course['title']} applied exercise (demo)",
            "body_text": f"Complete a short applied exercise based on the latest {course['title']} lecture and explain your reasoning. This assignment is synthetic demo data.",
            "origin": "synthetic",
            "source_ref": f"bb://{course['external_id']}/assignments/1",
            "posted_at": "2026-09-22T08:00:00Z",
            "due_at": "2026-10-05T20:59:00Z",
            "modified_at": "2026-09-22T08:00:00Z",
        },
    ]


def build_payload(root: Path, student_id: str) -> dict:
    courses: list[dict] = []
    for spec in COURSES:
        folder = root.joinpath(*spec["folder"])
        if not folder.is_dir():
            print(f"warning: approved folder not found: {folder}", file=sys.stderr)
            continue
        items = synthetic_items(spec)
        for path in eligible_files(folder, bool(spec["recursive"])):
            try:
                raw = path.read_bytes()
                text = extract_pdf(path) if path.suffix.lower() == ".pdf" else extract_pptx(path)
            except Exception as exc:
                print(f"warning: skipped {path.name}: {exc}", file=sys.stderr)
                continue
            if not text.strip():
                print(f"warning: no extractable text: {path.name}", file=sys.stderr)
                continue
            fingerprint = hashlib.sha256(raw).hexdigest()
            external_id = f"lecture-{fingerprint[:20]}"
            items.append({
                "external_id": external_id,
                "content_type": "lecture",
                "title": path.stem,
                "body_text": text,
                "filename": path.name,
                "mime_type": "application/pdf" if path.suffix.lower() == ".pdf" else "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "source_ref": f"bb://{spec['external_id']}/content/{external_id}",
                "origin": "local_material",
                "checksum": fingerprint,
                "posted_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
                "modified_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
            })
        courses.append({
            "external_id": spec["external_id"],
            "code": spec["code"],
            "title": spec["title"],
            "term": "Sixth Semester · Demo snapshot",
            "description": spec["description"],
            "items": items,
        })
    if not courses:
        raise RuntimeError("None of the approved course folders were found")
    return {"student_id": student_id, "courses": courses}


def upload(payload: dict, api_url: str, token: str) -> dict:
    request = Request(
        f"{api_url.rstrip('/')}/internal/demo/blackboard/import",
        data=json.dumps(payload).encode("utf-8"),
        method="POST",
        headers={"Content-Type": "application/json", "X-Farq-Internal-Token": token},
    )
    try:
        with urlopen(request, timeout=120) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        raise RuntimeError(f"Farq rejected the snapshot ({exc.code}): {exc.read().decode('utf-8', errors='replace')}") from exc
    except URLError as exc:
        raise RuntimeError(f"Farq API is unavailable: {exc.reason}") from exc


def main() -> None:
    parser = argparse.ArgumentParser(description="Import approved local lectures as a Blackboard demo snapshot")
    parser.add_argument("--root", required=True, type=Path, help="University folder containing the approved semester folders")
    parser.add_argument("--student", default="demo-student")
    parser.add_argument("--api", default="http://127.0.0.1:8000")
    parser.add_argument("--dry-run", action="store_true", help="Extract and summarize without uploading")
    args = parser.parse_args()
    root = args.root.resolve()
    payload = build_payload(root, args.student)
    summary = {course["code"]: len(course["items"]) for course in payload["courses"]}
    if args.dry_run:
        print(json.dumps({"student_id": args.student, "courses": summary}, ensure_ascii=False, indent=2))
        return
    env = read_env(Path(__file__).resolve().parents[1] / ".env")
    token = os.getenv("FARQ_INTERNAL_TOKEN") or env.get("FARQ_INTERNAL_TOKEN", "farq-internal-dev")
    result = upload(payload, args.api, token)
    print(json.dumps({**result, "content_by_course": summary}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
