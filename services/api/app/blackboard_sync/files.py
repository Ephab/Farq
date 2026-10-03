"""Attachment text for the Blackboard sync. Bytes live only in memory; only redacted text is kept."""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

from ..sources.pdf_text import redact

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_FILES = 60
MAX_CHARS = 40_000
MAX_PDF_PAGES = 80
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
TEXT_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".md"}
SYLLABUS = re.compile(r"syllabus|course (outline|spec|plan)|خطة المقرر|توصيف", re.I)


@dataclass(frozen=True)
class Attachment:
    key: str  # "{course_id}:{content_id}:{attachment_id}" — ingest looks texts up by this
    course_id: str
    content_id: str
    name: str
    url: str
    size: int | None


def _ext(name: str) -> str:
    return "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""


def select_attachments(export: dict) -> list[Attachment]:
    """Readable files, current courses first, syllabi first, newest first; capped."""
    current = {c.get("id") for c in export.get("courses", []) if c.get("is_current")}
    ranked: list[tuple[bool, bool, str, Attachment]] = []
    for item in export.get("content", []):
        for att in item.get("attachments") or []:
            name, url, size = str(att.get("name") or ""), att.get("download_url") or "", att.get("size")
            if not url or not att.get("id") or _ext(name) not in TEXT_EXTENSIONS:
                continue
            if isinstance(size, (int, float)) and size > MAX_FILE_BYTES:
                continue
            ranked.append((
                item.get("course_id") in current,
                bool(SYLLABUS.search(f"{item.get('title') or ''} {name}")),
                item.get("modified") or item.get("created") or "",
                Attachment(f"{item.get('course_id')}:{item.get('content_id')}:{att['id']}", item.get("course_id") or "",
                           item.get("content_id") or "", name, url, int(size) if isinstance(size, (int, float)) else None),
            ))
    ranked.sort(key=lambda row: (row[0], row[1], row[2]), reverse=True)
    return [row[3] for row in ranked[:MAX_FILES]]


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages[:MAX_PDF_PAGES])


def _pptx(data: bytes) -> str:
    from pptx import Presentation
    lines: list[str] = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        lines.append(f"[Slide {number}]")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
    return "\n".join(lines)


def _docx(data: bytes) -> str:
    from docx import Document
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


def _zip_ok(data: bytes) -> bool:
    """Office files are zips; refuse decompression bombs before a parser inflates them in memory."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            return len(infos) <= MAX_ZIP_ENTRIES and sum(i.file_size for i in infos) <= MAX_UNZIPPED_BYTES
    except zipfile.BadZipFile:
        return False


def extract_text(name: str, data: bytes) -> str | None:
    if not data or len(data) > MAX_FILE_BYTES:
        return None
    ext = _ext(name)
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext in {".pptx", ".docx"}:
            if not _zip_ok(data):
                return None
            text = _pptx(data) if ext == ".pptx" else _docx(data)
        elif ext in {".txt", ".md"}:
            text = data.decode("utf-8", errors="replace")
        else:
            return None
    except Exception:  # damaged, encrypted or mislabelled documents are skipped, never fatal
        return None
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 40:
        return None  # scanned PDF or an empty deck
    return redact(text)[:MAX_CHARS]
