from __future__ import annotations

"""Slide extension through the local Hermes gateway.

Mirrors ``app.quiz``: the browser never calls a model provider directly.
Two JSON-only prompts run on throwaway ``farq:slides-*`` sessions so deck
content never pollutes the coach's conversational memory. Nothing is written
to SQLite: slide source text is not an explicit student statement, so it must
not become a StudentFact, message, or proposal.

- ``run_suggest`` returns raw model text; the frontend parses
  ``{"topics": [...]}``.
- ``run_extend`` returns raw model text; the frontend parses
  ``{"slides": [...]}``.
- ``build_full_deck_pptx`` builds a downloadable .pptx locally (no model):
  the original slides are kept and the AI slides are appended on the same
  layouts with matching backgrounds and text styling. PDF originals are
  embedded as full-bleed page images first (rendered client-side).
"""

import base64
import binascii
import io
import time
import uuid

import httpx

from .hermes import (
    HERMES_URL,
    effective_hermes_key,
    raise_for_gateway_status,
    resolve_hermes_selection,
)

MAX_SOURCE_CHARS = 12_000
RUN_TIMEOUT_SECONDS = 180
POLL_INTERVAL_SECONDS = 2
MAX_ORIGINAL_BYTES = 25 * 1024 * 1024

SUGGEST_INSTRUCTIONS = " ".join([
    "You suggest follow-up topics that extend lecture slides.",
    "Do not call any tools. Return ONLY a JSON object: {\"topics\": [...]}. No markdown, no prose.",
    "Each topic: {\"id\":\"t1\",\"title\":\"...\",\"rationale\":\"one sentence, why it extends the deck\",\"related_slides\":\"Slide N or Page N\"}.",
    "Rules: titles are specific extension topics, not restatements of existing slides.",
    "Prefer gaps, next steps, real-world applications, and common misconceptions.",
    "Order most useful first. No content beyond the JSON object.",
])

EXTEND_INSTRUCTIONS = " ".join([
    "You write new lecture slides that extend an existing deck.",
    "Do not call any tools. Return ONLY a JSON object: {\"slides\": [...]}. No markdown, no prose.",
    "Each slide: {\"title\":\"...\",\"bullets\":[\"...\",\"...\"],\"speaker_notes\":\"one or two sentences\"}.",
    "Rules: 3-6 bullets per slide, each one concise sentence.",
    "You decide how many slides the topic needs within the requested length.",
    "Match the deck's terminology and depth; do not contradict the source.",
    "Ground new claims in the requested topic; keep speaker_notes brief.",
])

LENGTH_GUIDANCE = {
    "short": "brief extension: decide between 2-3 slides, only what the topic needs",
    "medium": "standard extension: decide between 4-6 slides, only what the topic needs",
    "long": "thorough extension: decide between 7-10 slides, only what the topic needs",
}


class SlidesRunError(RuntimeError):
    """Slide generation failure with the HTTP status the API should return."""

    def __init__(self, message: str, status: int = 502):
        super().__init__(message)
        self.status = status


def _gateway_key_or_raise(hermes_api_key: str | None) -> str:
    gateway_key = effective_hermes_key(hermes_api_key)
    if len(gateway_key) < 16:
        raise SlidesRunError(
            "Farq Hermes key is missing or too short; press Apply in Settings "
            "or set HERMES_API_KEY in the server .env",
            status=401,
        )
    return gateway_key


def build_suggest_input(source_text: str, count: int) -> str:
    source = source_text[:MAX_SOURCE_CHARS]
    return "\n".join([
        f"Suggest {count} extension topics for these slides.",
        "Each topic must be new material the student could add, not a summary of what exists.",
        "",
        "--- SLIDE TEXT START ---",
        source,
        "--- SLIDE TEXT END ---",
    ])


def build_extend_input(source_text: str, topic: str, length: str = "medium", design_hint: str = "") -> str:
    source = source_text[:MAX_SOURCE_CHARS]
    guidance = LENGTH_GUIDANCE.get(length, LENGTH_GUIDANCE["medium"])
    lines = [
        f"Extend the deck with new slides about: {topic.strip()}",
        f"Requested length: {length} ({guidance}).",
        "Decide the exact number of slides yourself — generate as many as the topic needs within that range, no more.",
        "Follow the deck's formatting habits: short titles, parallel bullet phrasing.",
    ]
    if design_hint.strip():
        lines.append(f"Deck design to match: {design_hint.strip()[:1500]}")
    lines += [
        "",
        "--- SLIDE TEXT START ---",
        source,
        "--- SLIDE TEXT END ---",
    ]
    return "\n".join(lines)


def _run_prompt(
    kind: str,
    prompt: str,
    provider: str | None,
    model: str | None,
    gateway_key: str,
) -> dict:
    try:
        run_model, run_provider = resolve_hermes_selection(provider, model)
    except ValueError as exc:
        raise SlidesRunError(str(exc), status=422) from exc
    session_id = f"slides-{kind}-{uuid.uuid4().hex[:12]}"
    headers = {
        "Authorization": f"Bearer {gateway_key}",
        "Idempotency-Key": f"slides-{session_id}",
        "X-Hermes-Session-Key": f"farq:slides:{session_id}",
    }
    instructions = SUGGEST_INSTRUCTIONS if kind == "suggest" else EXTEND_INSTRUCTIONS
    payload = {
        "input": prompt,
        "session_id": session_id,
        "instructions": instructions,
        "model": run_model,
        "provider": run_provider,
    }
    try:
        with httpx.Client(timeout=20) as client:
            try:
                response = client.post(f"{HERMES_URL}/v1/runs", headers=headers, json=payload)
                raise_for_gateway_status(response)
            except RuntimeError as exc:
                raise SlidesRunError(str(exc), status=401) from exc
            run_id = response.json()["run_id"]
            deadline = time.monotonic() + RUN_TIMEOUT_SECONDS
            while time.monotonic() < deadline:
                poll = client.get(f"{HERMES_URL}/v1/runs/{run_id}", headers=headers)
                raise_for_gateway_status(poll)
                state = poll.json()
                status = state.get("status")
                if status == "completed":
                    output = (state.get("output") or "").strip()
                    if not output:
                        raise SlidesRunError("Hermes returned an empty answer — try a smaller deck", status=502)
                    return {"output": output, "model": run_model, "provider": run_provider}
                if status in {"failed", "cancelled"}:
                    raise SlidesRunError(state.get("error") or f"Hermes run {status}", status=502)
                time.sleep(POLL_INTERVAL_SECONDS)
            raise SlidesRunError("Hermes did not finish within 180 seconds", status=504)
    except SlidesRunError:
        raise
    except httpx.HTTPError as exc:
        raise SlidesRunError(f"Hermes gateway unavailable: {exc}", status=502) from exc


def run_suggest(
    source_text: str,
    count: int = 5,
    provider: str | None = None,
    model: str | None = None,
    hermes_api_key: str | None = None,
) -> dict:
    gateway_key = _gateway_key_or_raise(hermes_api_key)
    return _run_prompt("suggest", build_suggest_input(source_text, count), provider, model, gateway_key)


def run_extend(
    source_text: str,
    topic: str,
    length: str = "medium",
    provider: str | None = None,
    model: str | None = None,
    hermes_api_key: str | None = None,
    design_hint: str = "",
) -> dict:
    if length not in LENGTH_GUIDANCE:
        raise SlidesRunError(f"Unknown extension length: {length}", status=422)
    gateway_key = _gateway_key_or_raise(hermes_api_key)
    return _run_prompt("extend", build_extend_input(source_text, topic, length, design_hint), provider, model, gateway_key)


def decode_original_pptx(original_pptx_base64: str | None) -> bytes | None:
    """Decode optional original .pptx bytes, enforcing the 25MB upload cap."""
    if not original_pptx_base64:
        return None
    try:
        raw = base64.b64decode(original_pptx_base64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise SlidesRunError("Original file is not valid base64.", status=422) from exc
    if len(raw) > MAX_ORIGINAL_BYTES:
        raise SlidesRunError("Original file is larger than 25MB.", status=422)
    if raw[:2] != b"PK":
        raise SlidesRunError("Original file is not a .pptx (expected a zip container).", status=422)
    return raw


MAX_IMAGE_BYTES = 10 * 1024 * 1024


def decode_image_list(images: list[str] | None) -> list[bytes]:
    """Decode rendered original pages (data URLs or raw base64) for PDF decks."""
    raw_list = images or []
    if len(raw_list) > 60:
        raise SlidesRunError("Too many original pages (max 60).", status=422)
    decoded: list[bytes] = []
    for item in raw_list:
        text = item.strip()
        if "," in text and "base64" in text.split(",", 1)[0]:
            text = text.split(",", 1)[1]
        try:
            raw = base64.b64decode(text, validate=True)
        except (binascii.Error, ValueError) as exc:
            raise SlidesRunError("An original page image is not valid base64.", status=422) from exc
        if len(raw) > MAX_IMAGE_BYTES:
            raise SlidesRunError("An original page image is larger than 10MB.", status=422)
        if raw[:8] != b"\x89PNG\r\n\x1a\n" and raw[:2] != b"\xff\xd8":
            raise SlidesRunError("Original page images must be PNG or JPEG.", status=422)
        decoded.append(raw)
    return decoded


def _solid_bg_rgb(slide):
    """Solid background color of a slide, or None (theme default/gradient/picture)."""
    try:
        fill = slide.background.fill
        if fill.type is not None and int(fill.type) == 1:  # MSO_FILL.SOLID
            return fill.fore_color.rgb
    except Exception:
        return None
    return None


def _run_style(run) -> dict:
    """Extract font styling from a run; missing values stay None (inherit)."""
    style: dict = {"name": None, "size": None, "bold": None, "color": None}
    try:
        font = run.font
    except Exception:
        return style
    try:
        style["name"] = font.name or None
    except Exception:
        pass
    try:
        style["size"] = font.size.pt if font.size is not None else None
    except Exception:
        pass
    try:
        style["bold"] = font.bold
    except Exception:
        pass
    try:
        style["color"] = font.color.rgb if font.color is not None else None
    except Exception:
        pass
    return style


def _shape_text_style(shape) -> dict | None:
    """Style of the first text run in a shape, or None when unavailable."""
    try:
        if not shape.has_text_frame:
            return None
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.text.strip():
                    return _run_style(run)
        return None
    except Exception:
        return None


def _apply_style(run, style: dict | None) -> None:
    if not style:
        return
    try:
        if style.get("name"):
            run.font.name = style["name"]
        if style.get("size"):
            from pptx.util import Pt
            run.font.size = Pt(style["size"])
        if style.get("bold") is not None:
            run.font.bold = style["bold"]
        if style.get("color") is not None:
            run.font.color.rgb = style["color"]
    except Exception:
        pass


def _reference_styling(prs, content_layout):
    """Background + title/body text styling sampled from the deck's own slides.

    Prefers the first slide already using the content layout so appended
    slides genuinely look like they belong to the deck.
    """
    bg = None
    title_style = None
    body_style = None
    reference = None
    for slide in prs.slides:
        try:
            if slide.slide_layout is content_layout:
                reference = slide
                break
        except Exception:
            continue
    if reference is None and len(prs.slides) > 0:
        reference = prs.slides[0]
    if reference is None:
        return bg, title_style, body_style
    bg = _solid_bg_rgb(reference)
    try:
        title_shape = reference.shapes.title
        if title_shape is not None:
            title_style = _shape_text_style(title_shape)
        for shape in reference.placeholders:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                body_style = _shape_text_style(shape)
                if body_style:
                    break
    except Exception:
        pass
    return bg, title_style, body_style


def _write_content_slide(prs, layout, bg, title_style, body_style, title: str, bullets: list[str], notes: str) -> None:
    from pptx.util import Pt

    new_slide = prs.slides.add_slide(layout)
    if bg is not None:
        try:
            new_slide.background.fill.solid()
            new_slide.background.fill.fore_color.rgb = bg
        except Exception:
            pass
    if new_slide.shapes.title is not None:
        new_slide.shapes.title.text = ""
        tf = new_slide.shapes.title.text_frame
        tf.clear()
        run = tf.paragraphs[0].add_run()
        run.text = title
        _apply_style(run, title_style)
    else:
        tx_box = new_slide.shapes.add_textbox(Pt(36), Pt(28), Pt(648), Pt(80))
        run = tx_box.text_frame.paragraphs[0].add_run()
        run.text = title
        _apply_style(run, title_style)
    body_written = False
    body = body_placeholder(new_slide)
    if body is not None:
        tf = body.text_frame
        tf.clear()
        for i, bullet in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.level = 0
            run = para.add_run()
            run.text = bullet
            _apply_style(run, body_style)
        body_written = True
    if not body_written:
        tx_box = new_slide.shapes.add_textbox(Pt(36), Pt(120), Pt(648), Pt(360))
        tf = tx_box.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.level = 0
            run = para.add_run()
            run.text = f"• {bullet}"
            _apply_style(run, body_style)
    if notes:
        try:
            new_slide.notes_slide.notes_text_frame.text = notes
        except Exception:
            pass


def build_full_deck_pptx(
    original_filename: str,
    topic: str,
    slides: list[dict],
    original_bytes: bytes | None = None,
    original_images: list[bytes] | None = None,
) -> bytes:
    """Build one downloadable .pptx: originals first, AI slides appended.

    - PPTX originals are kept intact; new slides reuse the deck's most-used
      content layout with its background and title/body text styling copied
      over, so they visually match the deck.
    - PDF originals arrive as client-rendered page images embedded full-bleed
      ahead of the new slides (default theme for the AI section).
    - A divider slide marks where the AI extension starts (provenance).
    """
    from pptx import Presentation
    from pptx.util import Pt

    images = original_images or []
    if original_bytes:
        prs = Presentation(io.BytesIO(original_bytes))
        original_count = len(prs.slides)
    elif images:
        prs = Presentation()
        blank = _blank_layout(prs)
        for raw in images:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(io.BytesIO(raw), Pt(0), Pt(0), width=prs.slide_width, height=prs.slide_height)
        original_count = len(prs.slides)
    else:
        prs = Presentation()
        original_count = 0

    def pick_layout():
        # Most-used layout with a body placeholder wins — that is the deck's
        # real content design, not just the template default.
        counts: dict = {}
        layouts: dict = {}
        for slide in prs.slides:
            try:
                layout = slide.slide_layout
            except Exception:
                continue
            key = id(layout)
            counts[key] = counts.get(key, 0) + 1
            layouts[key] = layout
        for key in sorted(counts, key=lambda k: -counts[k]):
            if _layout_has_body(layouts[key]):
                return layouts[key]
        if len(prs.slide_layouts) > 1:
            return prs.slide_layouts[1]
        return prs.slide_layouts[0]

    content_layout = pick_layout() if original_count else (
        prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    )
    title_layout = prs.slide_layouts[0] if len(prs.slide_layouts) > 0 else content_layout
    bg, title_style, body_style = _reference_styling(prs, content_layout) if original_count else (None, None, None)

    # Divider: provenance + topic, styled like the deck.
    _write_content_slide(
        prs, title_layout, bg, title_style, body_style,
        f"New: {topic.strip() or 'Extension'}",
        [f"The following {len(slides)} slide(s) extend {original_filename} — generated with Hermes, review before presenting."],
        "",
    )

    for slide in slides:
        title = str(slide.get("title", "")).strip() or "New slide"
        bullets = [str(b).strip() for b in slide.get("bullets", []) if str(b).strip()][:8]
        notes = str(slide.get("speaker_notes", "") or "").strip()
        _write_content_slide(prs, content_layout, bg, title_style, body_style, title, bullets, notes)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def body_placeholder(slide):
    """First text placeholder that is not the title (slide or layout)."""
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None
    try:
        for shape in slide.placeholders:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                return shape
    except Exception:
        return None
    return None


def _layout_has_body(layout) -> bool:
    try:
        for ph in layout.placeholders:
            if not ph.has_text_frame:
                continue
            try:
                is_title = ph.placeholder_format.type in (1, 3)
            except Exception:
                is_title = False
            if not is_title:
                return True
        return False
    except Exception:
        return False


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        try:
            if len(list(layout.placeholders)) == 0:
                return layout
        except Exception:
            continue
    if len(prs.slide_layouts) > 6:
        return prs.slide_layouts[6]
    return prs.slide_layouts[0]


# Backwards-compatible alias (same full-deck behavior, no originals).
def build_extended_pptx(
    original_filename: str,
    topic: str,
    slides: list[dict],
    original_bytes: bytes | None = None,
) -> bytes:
    return build_full_deck_pptx(original_filename, topic, slides, original_bytes, None)
