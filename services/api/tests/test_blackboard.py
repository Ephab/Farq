import importlib.util
import os
import tempfile
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation


TEST_DB = Path(tempfile.gettempdir()) / f"farq-blackboard-{uuid.uuid4()}.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB.as_posix()}"

from app.database import engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Student  # noqa: E402


INTERNAL = {"X-Farq-Internal-Token": "farq-internal-dev"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client
    engine.dispose()
    TEST_DB.unlink(missing_ok=True)


def snapshot_payload() -> dict:
    return {
        "student_id": "demo-student",
        "courses": [{
            "external_id": "ARTI-404",
            "code": "ARTI 404",
            "title": "Machine Learning",
            "term": "Sixth Semester",
            "description": "Machine-learning foundations.",
            "items": [
                {
                    "external_id": "lecture-logistic",
                    "content_type": "lecture",
                    "title": "Logistic Regression",
                    "body_text": "Logistic regression uses the sigmoid function for probabilistic binary classification. " * 300,
                    "filename": "Week 8.pdf",
                    "mime_type": "application/pdf",
                    "source_ref": "bb://ARTI-404/content/lecture-logistic",
                    "origin": "local_material",
                    "modified_at": "2026-09-24T09:00:00Z",
                },
                {
                    "external_id": "announcement-1",
                    "content_type": "announcement",
                    "title": "New material (demo)",
                    "body_text": "Synthetic Blackboard announcement.",
                    "origin": "synthetic",
                    "modified_at": "2026-09-25T09:00:00Z",
                },
            ],
        }],
    }


def test_blackboard_import_requires_internal_token(client: TestClient):
    assert client.post("/internal/demo/blackboard/import", json=snapshot_payload()).status_code == 401


def test_import_and_every_read_tool_endpoint(client: TestClient):
    imported = client.post("/internal/demo/blackboard/import", json=snapshot_payload(), headers=INTERNAL)
    assert imported.status_code == 200, imported.text
    assert imported.json()["items"] == 2

    status = client.get("/api/students/demo-student/blackboard/status").json()
    assert status == {"connected": True, "mode": "preindexed_demo", "read_only": True, "courses": 1, "last_synced_at": status["last_synced_at"]}

    courses = client.get("/internal/hermes/students/demo-student/blackboard/courses", headers=INTERNAL)
    assert courses.status_code == 200
    course = courses.json()["courses"][0]
    assert course["content_count"] == 2

    listed = client.get(f"/internal/hermes/students/demo-student/blackboard/courses/{course['id']}/content?content_type=lecture", headers=INTERNAL)
    assert listed.status_code == 200
    item = listed.json()["items"][0]
    assert "body_text" not in item

    searched = client.get("/internal/hermes/students/demo-student/blackboard/search?query=sigmoid", headers=INTERNAL)
    assert searched.status_code == 200
    assert searched.json()["results"][0]["id"] == item["id"]
    assert len(searched.json()["results"][0]["snippet"]) <= 360

    first = client.get(f"/internal/hermes/students/demo-student/blackboard/items/{item['id']}", headers=INTERNAL).json()
    assert len(first["text"]) <= 12_000
    assert first["next_cursor"] == 12_000
    second = client.get(f"/internal/hermes/students/demo-student/blackboard/items/{item['id']}?cursor={first['next_cursor']}", headers=INTERNAL)
    assert second.status_code == 200

    updates = client.get("/internal/hermes/students/demo-student/blackboard/updates?since=2026-09-25T00:00:00Z", headers=INTERNAL)
    assert updates.status_code == 200
    assert updates.json()["items"][0]["origin"] == "synthetic"


def test_blackboard_student_isolation(client: TestClient):
    from app.database import SessionLocal
    with SessionLocal() as db:
        db.add(Student(id="other-student", display_name="Other Student"))
        db.commit()
    courses = client.get("/internal/hermes/students/demo-student/blackboard/courses", headers=INTERNAL).json()["courses"]
    course_id = courses[0]["id"]
    item_id = client.get(f"/internal/hermes/students/demo-student/blackboard/courses/{course_id}/content", headers=INTERNAL).json()["items"][0]["id"]
    assert client.get(f"/internal/hermes/students/other-student/blackboard/courses/{course_id}/content", headers=INTERNAL).status_code == 404
    assert client.get(f"/internal/hermes/students/other-student/blackboard/items/{item_id}", headers=INTERNAL).status_code == 404


def test_blackboard_import_is_idempotent(client: TestClient):
    assert client.post("/internal/demo/blackboard/import", json=snapshot_payload(), headers=INTERNAL).status_code == 200
    courses = client.get("/internal/hermes/students/demo-student/blackboard/courses", headers=INTERNAL).json()["courses"]
    assert len(courses) == 1
    assert courses[0]["content_count"] == 2


def test_blackboard_import_replaces_stale_snapshot_items(client: TestClient):
    payload = snapshot_payload()
    payload["courses"][0]["items"] = payload["courses"][0]["items"][:1]
    assert client.post("/internal/demo/blackboard/import", json=payload, headers=INTERNAL).status_code == 200
    course = client.get("/internal/hermes/students/demo-student/blackboard/courses", headers=INTERNAL).json()["courses"][0]
    assert course["content_count"] == 1
    search = client.get("/internal/hermes/students/demo-student/blackboard/search?query=synthetic", headers=INTERNAL)
    assert search.json()["results"] == []


def test_demo_importer_extracts_pptx_and_filters_blocked_names(tmp_path: Path):
    repo = Path(__file__).resolve().parents[3]
    script = repo / "scripts" / "import_blackboard_demo.py"
    spec = importlib.util.spec_from_file_location("import_blackboard_demo", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Gradient Descent"
    slide.placeholders[1].text = "Follow the negative gradient of the loss."
    lecture = tmp_path / "Lecture 1.pptx"
    deck.save(lecture)
    blocked = tmp_path / "Quiz solutions.pptx"
    deck.save(blocked)

    assert "negative gradient" in module.extract_pptx(lecture)
    assert module.eligible_files(tmp_path, True) == [lecture]
