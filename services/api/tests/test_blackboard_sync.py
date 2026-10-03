import copy
import json
import os
import tempfile
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

TEST_DB = Path(tempfile.gettempdir()) / f"waypoint-bbsync-{uuid.uuid4()}.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB.as_posix()}"
os.environ["BLACKBOARD_SYNC_ENABLED"] = "false"

from sqlalchemy import select  # noqa: E402

from app.database import SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import BlackboardContentItem, BlackboardCourse, BlackboardGrade, EvidenceItem, Student  # noqa: E402
from app.blackboard_sync import ingest  # noqa: E402

FIXTURE = Path(__file__).parent / "fixtures" / "blackboard-export-sample.json"
SYLLABUS_KEY = "_101_1:_file1:_att9"


def sample() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client
    engine.dispose()
    try:
        TEST_DB.unlink(missing_ok=True)
    except PermissionError:
        pass  # store_evidence's background observer may still hold the file on Windows


@pytest.fixture()
def student(client):
    sid = f"bb-{uuid.uuid4().hex[:8]}"
    db = SessionLocal()
    db.add(Student(id=sid, display_name="BB Student"))
    db.commit()
    db.close()
    return sid


def ingest_sample(sid: str, export: dict | None = None, texts: dict | None = None):
    db = SessionLocal()
    try:
        result = ingest.ingest_export(db, sid, export or sample(), texts or {})
        db.commit()
        return result
    finally:
        db.close()


def rows(model, **where):
    db = SessionLocal()
    try:
        query = select(model)
        for key, value in where.items():
            query = query.where(getattr(model, key) == value)
        return db.scalars(query).all()
    finally:
        db.close()


def test_plain_strips_html_and_entities():
    assert ingest.plain("<h5>QUIZ</h5><p>A &amp; B</p>") == "QUIZ\nA & B"
    assert ingest.plain(None) == ""


def test_course_label_drops_section_suffix():
    assert ingest.course_label("Deep Learning-7MA1") == "Deep Learning"
    assert ingest.course_label("Machine learning-6MA2") == "Machine learning"
    assert ingest.course_label("Ethics") == "Ethics"


def test_ingest_maps_courses_items_grades(student):
    result = ingest_sample(student, texts={SYLLABUS_KEY: "Midterm covers chapters 1-4."})
    courses = {c.external_id: c for c in rows(BlackboardCourse, student_id=student)}
    assert set(courses) == {"_101_1", "_102_1", "_090_1"}
    dl = courses["_101_1"]
    assert dl.source_kind == "blackboard_live" and dl.is_current and dl.term == "Fall 2026"
    assert json.loads(dl.instructors_json)[0]["name"] == "Sara Ali"
    assert json.loads(dl.grade_summary_json)["percentage"] == 90
    items = {i.external_id: i for i in rows(BlackboardContentItem, course_id=dl.id)}
    assert items["ann:_a1"].content_type == "announcement" and "Topics: Backprop" in items["ann:_a1"].body_text
    assert items["asmt:_col4"].content_type == "assignment" and items["asmt:_col4"].due_at is not None
    assert items["content:_file1"].content_type == "syllabus"
    assert "Midterm covers chapters 1-4." in items["content:_file1"].body_text
    assert items["content:_c1"].content_type == "lecture" and items["content:_c1"].body_text == "Hello & welcome"
    assert "content:_f1" not in items  # empty folder rows carry nothing useful
    grade = rows(BlackboardGrade, course_id=dl.id)[0]
    assert grade.score == 9 and grade.possible == 10 and grade.feedback == "Good work"
    assert result.courses == 3 and result.current_courses == 2
    assert result.upcoming_deadlines == 1 and result.overdue == 1 and result.files_read == 1


def test_ingest_suggests_course_evidence_once(student):
    first = ingest_sample(student)
    evidence = rows(EvidenceItem, student_id=student)
    titles = sorted(e.title for e in evidence)
    # Two sections of Deep Learning collapse to one; the completed course carries its grade.
    assert titles == ["Deep Learning", "Machine learning"]
    assert all(e.status == "suggested" for e in evidence)
    completed = next(e for e in evidence if e.title == "Machine learning")
    assert json.loads(completed.data_json)["grade"] == "B+"
    assert first.new_evidence == 2
    assert ingest_sample(student).new_evidence == 0


def test_dismissed_evidence_not_resurrected(student):
    ingest_sample(student)
    db = SessionLocal()
    for e in db.scalars(select(EvidenceItem).where(EvidenceItem.student_id == student)).all():
        e.status = "dismissed"
    db.commit()
    db.close()
    ingest_sample(student)
    assert {e.status for e in rows(EvidenceItem, student_id=student)} == {"dismissed"}


def test_first_live_sync_replaces_demo_rows(student):
    db = SessionLocal()
    demo = BlackboardCourse(student_id=student, external_id="ARTI-404", title="Demo", source_kind="blackboard_demo")
    db.add(demo)
    db.commit()
    db.close()
    ingest_sample(student)
    assert "ARTI-404" not in {c.external_id for c in rows(BlackboardCourse, student_id=student)}


def test_empty_export_keeps_existing_rows(student):
    ingest_sample(student)
    empty = sample()
    empty["courses"] = []
    with pytest.raises(ValueError):
        ingest_sample(student, empty)
    assert len(rows(BlackboardCourse, student_id=student)) == 3


def test_failed_source_keeps_old_items(student):
    ingest_sample(student)
    later = sample()
    later["announcements"] = []
    later["diagnostics"]["sources"] = [s for s in later["diagnostics"]["sources"] if s["source"] != "announcements:_101_1"]
    later["diagnostics"]["sources"].append({"source": "announcements:_101_1", "status": "http_500"})
    ingest_sample(student, later)
    course = next(c for c in rows(BlackboardCourse, student_id=student) if c.external_id == "_101_1")
    assert any(i.external_id == "ann:_a1" for i in rows(BlackboardContentItem, course_id=course.id))


def test_removed_item_deleted_when_source_ok(student):
    ingest_sample(student)
    later = sample()
    later["announcements"] = []
    ingest_sample(student, later)
    course = next(c for c in rows(BlackboardCourse, student_id=student) if c.external_id == "_101_1")
    assert not any(i.external_id == "ann:_a1" for i in rows(BlackboardContentItem, course_id=course.id))


def test_unchanged_item_keeps_modified_at(student):
    ingest_sample(student)
    course = next(c for c in rows(BlackboardCourse, student_id=student) if c.external_id == "_101_1")
    before = {i.external_id: i.modified_at for i in rows(BlackboardContentItem, course_id=course.id)}
    ingest_sample(student)
    after = {i.external_id: i.modified_at for i in rows(BlackboardContentItem, course_id=course.id)}
    assert before["asmt:_col4"] == after["asmt:_col4"]



def test_unread_attachment_keeps_earlier_text_until_removed(student):
    def syllabus_body():
        course = next(c for c in rows(BlackboardCourse, student_id=student) if c.external_id == "_101_1")
        item = next(i for i in rows(BlackboardContentItem, course_id=course.id) if i.external_id == "content:_file1")
        return item.body_text

    ingest_sample(student, texts={SYLLABUS_KEY: "Midterm covers chapters 1-4."})
    ingest_sample(student, texts={})
    assert "Midterm covers chapters 1-4." in syllabus_body()
    removed = sample()
    next(c for c in removed["content"] if c["content_id"] == "_file1")["attachments"] = []
    ingest_sample(student, removed, texts={})
    assert "Midterm covers chapters 1-4." not in syllabus_body()


import io  # noqa: E402

from cryptography.fernet import Fernet  # noqa: E402

from app.blackboard_sync import credentials, files  # noqa: E402
from app.models import BlackboardConnection  # noqa: E402


@pytest.fixture()
def fernet_key(monkeypatch):
    monkeypatch.setenv("WAYPOINT_TOKEN_ENCRYPTION_KEY", Fernet.generate_key().decode())


def test_password_and_session_are_sealed(fernet_key):
    conn = BlackboardConnection(student_id="s1", username="2240000000")
    credentials.remember_password(conn, "hunter2-secret")
    credentials.save_session(conn, {"cookies": [{"name": "BbRouter", "value": "abc"}]})
    assert "hunter2-secret" not in (conn.password_enc or "")
    assert "BbRouter" not in (conn.session_enc or "")
    assert credentials.saved_password(conn) == "hunter2-secret"
    assert credentials.saved_session(conn)["cookies"][0]["value"] == "abc"
    credentials.forget_password(conn)
    credentials.clear_session(conn)
    assert credentials.saved_password(conn) is None and credentials.saved_session(conn) is None


def test_without_key_password_is_not_remembered(monkeypatch):
    monkeypatch.setenv("WAYPOINT_TOKEN_ENCRYPTION_KEY", "")
    assert credentials.can_remember() is False
    conn = BlackboardConnection(student_id="s-nokey", username="u")
    credentials.save_session(conn, {"cookies": []})
    assert conn.session_enc is None and credentials.saved_session(conn) == {"cookies": []}  # memory only
    credentials.clear_session(conn)


def test_rotated_key_reads_as_not_saved(fernet_key, monkeypatch):
    conn = BlackboardConnection(student_id="s2", username="u")
    credentials.remember_password(conn, "pw")
    monkeypatch.setenv("WAYPOINT_TOKEN_ENCRYPTION_KEY", Fernet.generate_key().decode())
    assert credentials.saved_password(conn) is None


def test_select_attachments_prefers_current_syllabus_and_caps():
    export = sample()
    picked = files.select_attachments(export)
    assert [a.key for a in picked] == [SYLLABUS_KEY]
    assert picked[0].url.endswith("/attachments/_att9/download")
    many = copy.deepcopy(export)
    base = many["content"][1]
    many["content"] = [dict(base, content_id=f"_x{i}", attachments=[dict(base["attachments"][0], id=f"_a{i}", name=f"f{i}.pdf")]) for i in range(90)]
    many["content"].append(dict(base, content_id="_big", attachments=[dict(base["attachments"][0], id="_big", size=files.MAX_FILE_BYTES + 1)]))
    many["content"].append(dict(base, content_id="_zip", attachments=[dict(base["attachments"][0], id="_zip", name="code.zip")]))
    chosen = files.select_attachments(many)
    assert len(chosen) == files.MAX_FILES
    assert not any(a.key.endswith(":_big") or a.key.endswith(":_zip") for a in chosen)


def test_extract_text_from_office_files_redacts_ids():
    from docx import Document
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("Midterm covers chapters 1-4. Contact 2240003321 for questions about grading policy.")
    buffer = io.BytesIO()
    doc.save(buffer)
    text = files.extract_text("Syllabus.docx", buffer.getvalue())
    assert "Midterm covers chapters 1-4." in text and "2240003321" not in text

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Backpropagation"
    slide.placeholders[1].text = "Chain rule applied layer by layer through the network graph."
    buffer = io.BytesIO()
    deck.save(buffer)
    assert "Chain rule" in files.extract_text("Lecture 3.pptx", buffer.getvalue())


def test_extract_text_rejects_garbage_and_oversize():
    assert files.extract_text("broken.pdf", b"%PDF-not-really") is None
    assert files.extract_text("archive.zip", b"PK...") is None
    assert files.extract_text("huge.txt", b"a" * (files.MAX_FILE_BYTES + 1)) is None
    long = files.extract_text("notes.txt", ("word " * 20_000).encode())
    assert long is not None and len(long) <= files.MAX_CHARS


def test_extract_text_rejects_zip_bombs():
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (101 * 1024 * 1024))
    assert len(buffer.getvalue()) < files.MAX_FILE_BYTES
    assert files.extract_text("bomb.docx", buffer.getvalue()) is None
    assert files.extract_text("bomb.pptx", buffer.getvalue()) is None


import logging  # noqa: E402

from app.blackboard_sync import worker  # noqa: E402
from app.blackboard_sync.browser import BrowserResult, ExtractFailure, LoginFailure  # noqa: E402

SECRET = "S3cret-Pa55-word!"


class FakeBrowser:
    """Stands in for Chromium: plays back an outcome and records what it was asked."""
    calls: list[dict] = []
    outcome: object = None

    def run(self, *, username, password, session_state, pick_attachments, progress):
        FakeBrowser.calls.append({"username": username, "password": password, "session": session_state})
        progress("extracting", "Courses: 3")
        if isinstance(FakeBrowser.outcome, Exception):
            raise FakeBrowser.outcome
        export = sample()
        picks = pick_attachments(export)
        return BrowserResult(export=export, session_state={"cookies": [{"name": "BbRouter", "value": "x"}]},
                             # Text bytes under a .txt name: extract_text dispatches on the extension.
                             files={a.key: ("syllabus.txt", b"Syllabus: midterm covers chapters 1-4 and the final covers all.") for a in picks})


@pytest.fixture()
def fake_browser(monkeypatch, fernet_key):
    FakeBrowser.calls, FakeBrowser.outcome = [], None
    monkeypatch.setattr(worker, "browser_factory", FakeBrowser)
    real_start = worker.start
    monkeypatch.setattr(worker, "start", lambda sid, pw, remember, background=True: real_start(sid, pw, remember, background=False))
    return FakeBrowser


def sync(client, sid, **body):
    return client.post(f"/api/students/{sid}/blackboard/sync", json=body or None)


def test_first_sync_with_credentials(client, student, fake_browser):
    response = sync(client, student, username="2240000000", password=SECRET, remember=True)
    assert response.status_code == 202
    status = client.get(f"/api/students/{student}/blackboard/sync").json()
    assert status["status"] == "done" and status["has_saved_login"] is True
    assert status["summary"]["current_courses"] == 2 and status["summary"]["files_read"] == 1
    assert status["next_sync_at"] is not None
    assert fake_browser.calls[0]["password"] == SECRET


def test_resync_uses_saved_session_and_password(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET)
    assert sync(client, student).status_code == 202
    second = fake_browser.calls[-1]
    assert second["password"] == SECRET and second["session"]["cookies"][0]["name"] == "BbRouter"


def test_remember_false_drops_password_keeps_session(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET, remember=False)
    status = client.get(f"/api/students/{student}/blackboard/sync").json()
    assert status["has_saved_login"] is False and status["status"] == "done"
    sync(client, student)
    assert fake_browser.calls[-1]["password"] is None and fake_browser.calls[-1]["session"]


def test_password_never_in_responses_or_logs(client, student, fake_browser, caplog):
    caplog.set_level(logging.DEBUG)
    bodies = [sync(client, student, username="2240000000", password=SECRET, remember=True).text,
              client.get(f"/api/students/{student}/blackboard/sync").text]
    fake_browser.outcome = LoginFailure("bad_password")
    bodies.append(sync(client, student, username="2240000000", password=SECRET).text)
    bodies.append(client.get(f"/api/students/{student}/blackboard/sync").text)
    assert all(SECRET not in body for body in bodies)
    assert SECRET not in caplog.text


def test_overlong_password_not_echoed(client, student, fake_browser):
    response = sync(client, student, username="u", password=SECRET * 40)
    assert response.status_code == 422 and SECRET not in response.text


def test_saved_password_rejected_is_wiped_and_loop_stops(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET)
    fake_browser.outcome = LoginFailure("bad_password")
    sync(client, student)  # periodic-style resync with the saved password
    status = client.get(f"/api/students/{student}/blackboard/sync").json()
    assert status["status"] == "failed" and status["failure_reason"] == "bad_password"
    assert status["has_saved_login"] is False and status["next_sync_at"] is None
    db = SessionLocal()
    assert student not in worker.due_students(db)
    db.close()


def test_three_typed_bad_passwords_wipe_everything(client, student, fake_browser):
    fake_browser.outcome = LoginFailure("bad_password")
    for _ in range(3):
        sync(client, student, username="2240000000", password="nope-nope")
    db = SessionLocal()
    conn = db.get(BlackboardConnection, student)
    assert conn.failed_logins == 3 and conn.password_enc is None and conn.session_enc is None
    db.close()


def test_unreachable_keeps_data_and_retries_later(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET)
    fake_browser.outcome = LoginFailure("unreachable")
    sync(client, student)
    status = client.get(f"/api/students/{student}/blackboard/sync").json()
    assert status["failure_reason"] == "unreachable" and status["next_sync_at"] is not None
    assert len(rows(BlackboardCourse, student_id=student)) == 3


def test_extract_failure_is_reported(client, student, fake_browser):
    fake_browser.outcome = ExtractFailure("Not logged into Blackboard")
    sync(client, student, username="2240000000", password=SECRET)
    assert client.get(f"/api/students/{student}/blackboard/sync").json()["failure_reason"] == "extract_failed"


def test_first_sync_needs_credentials(client, student, fake_browser):
    assert sync(client, student).status_code == 422
    assert sync(client, student, username="2240000000").status_code == 422


def test_second_post_while_running_is_single_flight(client, student, fake_browser, monkeypatch):
    sync(client, student, username="2240000000", password=SECRET)
    monkeypatch.setattr(worker, "_running", {student})
    before = len(fake_browser.calls)
    response = sync(client, student)
    assert response.status_code == 202 and len(fake_browser.calls) == before


def test_connector_switch_blocks_sync(client, student, fake_browser):
    from app.student_memory import set_connector
    db = SessionLocal()
    set_connector(db, student, "blackboard", False)
    db.commit()
    db.close()
    assert sync(client, student, username="2240000000", password=SECRET).status_code == 409


def test_forget_wipes_login(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET)
    assert client.delete(f"/api/students/{student}/blackboard/connection").json() == {"forgotten": True}
    status = client.get(f"/api/students/{student}/blackboard/sync").json()
    assert status["connected"] is False and status["has_saved_login"] is False
    assert len(rows(BlackboardCourse, student_id=student)) == 3  # synced course data stays


def test_other_student_cannot_read_status(client, student, fake_browser):
    other = client.post("/api/students", json={"display_name": "Intruder"}).json()["student_id"]
    response = client.get(f"/api/students/{student}/blackboard/sync", headers={"X-Waypoint-User": other, "x-test-no-auto": "1"})
    assert response.status_code in {403, 404}


def test_deadlines_endpoint_lists_upcoming_current(client, student, fake_browser):
    sync(client, student, username="2240000000", password=SECRET)
    items = client.get(f"/api/students/{student}/blackboard/deadlines").json()["items"]
    assert [i["title"] for i in items] == ["Project report"]
    assert items[0]["course"] == "Deep Learning-7MA1" and items[0]["overdue"] is False


def test_reset_interrupted_marks_running_rows_failed(client, student):
    db = SessionLocal()
    db.add(BlackboardConnection(student_id=student, username="u", status="extracting"))
    db.commit()
    worker.reset_interrupted(db)
    assert db.get(BlackboardConnection, student).failure_reason == "interrupted"
    db.close()
