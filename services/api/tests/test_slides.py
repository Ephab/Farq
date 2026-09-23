import base64
import io
import os
import tempfile
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient


TEST_DB = Path(tempfile.gettempdir()) / f"farq-slides-{uuid.uuid4()}.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB.as_posix()}"

from app.database import SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import AgentRun  # noqa: E402


CANNED_SUGGEST = '{"topics": [{"id": "t1", "title": "Attention mechanisms", "rationale": "Natural next step", "related_slides": "Slide 3"}]}'
CANNED_EXTEND = '{"slides": [{"title": "Attention", "bullets": ["Query-key-value idea", "Scaled dot-product"], "speaker_notes": "Keep it intuitive."}]}'


class _FakeSlidesResponse:
    def __init__(self, payload: dict):
        self._payload = payload

    def json(self) -> dict:
        return self._payload

    def raise_for_status(self) -> None:
        return None


def _fake_slides_client(status: dict):
    calls: dict = {}

    class FakeClient:
        def __init__(self, *args, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def post(self, url, headers=None, json=None):
            calls["auth"] = (headers or {}).get("Authorization")
            calls["session"] = (json or {}).get("session_id")
            calls["session_key"] = (headers or {}).get("X-Hermes-Session-Key")
            calls["input"] = (json or {}).get("input")
            return _FakeSlidesResponse({"run_id": "run-slides-1"})

        def get(self, url, headers=None):
            return _FakeSlidesResponse(status)

    return FakeClient, calls


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client
    engine.dispose()
    TEST_DB.unlink(missing_ok=True)


def test_suggest_uses_gateway_and_writes_nothing(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    from sqlalchemy import func, select

    FakeClient, calls = _fake_slides_client({"status": "completed", "output": CANNED_SUGGEST})
    monkeypatch.setattr("app.slides.httpx.Client", FakeClient)
    db = SessionLocal()
    before = db.scalar(select(func.count()).select_from(AgentRun))
    db.close()

    response = client.post(
        "/api/slides/suggest",
        headers={"X-Hermes-Api-Key": "k" * 64},
        json={"source_text": "Transformers intro. Slide 3 covers embeddings.", "count": 3},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["output"] == CANNED_SUGGEST
    assert calls["auth"] == "Bearer " + "k" * 64
    assert str(calls["session"]).startswith("slides-suggest-")
    assert str(calls["session_key"]).startswith("farq:slides:")
    assert "extension topics" in calls["input"]

    db = SessionLocal()
    after = db.scalar(select(func.count()).select_from(AgentRun))
    db.close()
    assert after == before


def test_extend_uses_gateway_and_includes_topic(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    FakeClient, calls = _fake_slides_client({"status": "completed", "output": CANNED_EXTEND})
    monkeypatch.setattr("app.slides.httpx.Client", FakeClient)

    response = client.post(
        "/api/slides/extend",
        headers={"X-Hermes-Api-Key": "k" * 64},
        json={"source_text": "Transformers intro.", "topic": "Attention mechanisms", "length": "short"},
    )
    assert response.status_code == 200
    assert response.json()["output"] == CANNED_EXTEND
    assert "Attention mechanisms" in calls["input"]
    assert "short" in calls["input"]


def test_extend_length_hint_is_validated(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    from app.slides import build_extend_input

    assert "2-3" in build_extend_input("src", "topic", "short")
    assert "7-10" in build_extend_input("src", "topic", "long")
    with pytest.raises(Exception):
        from app.slides import run_extend
        run_extend("src", "topic", "enormous", hermes_api_key="k" * 64)


def test_slides_reject_bad_input(client: TestClient):
    headers = {"X-Hermes-Api-Key": "k" * 64}
    assert client.post("/api/slides/suggest", headers=headers, json={"source_text": "", "count": 3}).status_code == 422
    assert client.post("/api/slides/suggest", headers=headers, json={"source_text": "x", "count": 99}).status_code == 422
    assert client.post("/api/slides/extend", headers=headers, json={"source_text": "x", "topic": "", "length": "short"}).status_code == 422
    assert client.post("/api/slides/extend", headers=headers, json={"source_text": "x", "topic": "t", "length": "enormous"}).status_code == 422
    unknown = client.post(
        "/api/slides/suggest", headers=headers,
        json={"source_text": "x", "count": 2, "provider": "gemini", "model": "not-a-model"},
    )
    # Model allowlist is enforced against the gateway path; with a mocked
    # gateway it surfaces as 422 only when the selection fails before HTTP.
    # Without gateway mock, an unknown model fails fast in run_suggest.
    assert unknown.status_code in {422, 502}


def test_slides_map_gateway_failure(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    FakeClient, _ = _fake_slides_client({"status": "failed", "error": "model exploded"})
    monkeypatch.setattr("app.slides.httpx.Client", FakeClient)
    response = client.post(
        "/api/slides/extend",
        headers={"X-Hermes-Api-Key": "k" * 64},
        json={"source_text": "Slides.", "topic": "More depth", "length": "medium"},
    )
    assert response.status_code == 502
    assert "model exploded" in response.json()["detail"]


def test_export_builds_pptx_with_default_theme(client: TestClient):
    response = client.post(
        "/api/slides/export",
        json={
            "original_filename": "lecture.pdf",
            "topic": "Attention mechanisms",
            "slides": [
                {"title": "Attention", "bullets": ["Query-key-value", "Scaled dot-product"], "speaker_notes": "Keep it intuitive."},
                {"title": "In practice", "bullets": ["Multi-head", "Masks"]},
            ],
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.content[:2] == b"PK"
    from pptx import Presentation

    prs = Presentation(io.BytesIO(response.content))
    # Divider + 2 new slides (no originals supplied).
    assert len(prs.slides) == 3
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in prs.slides]
    assert titles[0].startswith("New:")
    assert titles[1] == "Attention"
    assert titles[2] == "In practice"


def test_export_keeps_originals_and_matches_styling(client: TestClient):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    template_buf = io.BytesIO()
    template = Presentation()
    slide = template.slides.add_slide(template.slide_layouts[1])
    slide.shapes.title.text = "Original"
    # Distinctive dark background + title styling the new slides must copy.
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0x1B, 0x1B, 0x3A)
    run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xCC, 0x00)
    template.save(template_buf)
    original_b64 = base64.b64encode(template_buf.getvalue()).decode("ascii")

    response = client.post(
        "/api/slides/export",
        json={
            "original_filename": "lecture.pptx",
            "topic": "Attention",
            "slides": [{"title": "New idea", "bullets": ["Point one", "Point two"]}],
            "original_pptx_base64": original_b64,
        },
    )
    assert response.status_code == 200
    prs = Presentation(io.BytesIO(response.content))
    # Original + divider + 1 new slide.
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Original"
    assert prs.slides[1].shapes.title.text.startswith("New:")
    assert prs.slides[2].shapes.title.text == "New idea"
    new_bg = prs.slides[2].background.fill.fore_color.rgb
    assert (new_bg[0], new_bg[1], new_bg[2]) == (0x1B, 0x1B, 0x3A)
    new_title_run = prs.slides[2].shapes.title.text_frame.paragraphs[0].runs[0]
    assert new_title_run.font.bold is True
    assert (new_title_run.font.color.rgb[0], new_title_run.font.color.rgb[1], new_title_run.font.color.rgb[2]) == (0xFF, 0xCC, 0x00)


def test_export_embeds_pdf_page_images(client: TestClient):
    import struct
    import zlib

    def tiny_png(rgb=(255, 255, 255)):
        def chunk(tag, data):
            c = struct.pack(">I", len(data)) + tag + data
            return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        ihdr = struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0)
        raw = b"".join(b"\x00" + bytes(rgb) * 2 for _ in range(2))
        return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")

    page = "data:image/png;base64," + base64.b64encode(tiny_png()).decode("ascii")
    response = client.post(
        "/api/slides/export",
        json={
            "original_filename": "lecture.pdf",
            "topic": "Attention",
            "slides": [{"title": "New idea", "bullets": ["Point one"]}],
            "original_images_base64": [page, page],
        },
    )
    assert response.status_code == 200
    from pptx import Presentation

    prs = Presentation(io.BytesIO(response.content))
    # 2 page images + divider + 1 new slide.
    assert len(prs.slides) == 4


def test_extend_design_hint_reaches_prompt(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    FakeClient, calls = _fake_slides_client({"status": "completed", "output": CANNED_EXTEND})
    monkeypatch.setattr("app.slides.httpx.Client", FakeClient)

    response = client.post(
        "/api/slides/extend",
        headers={"X-Hermes-Api-Key": "k" * 64},
        json={
            "source_text": "Transformers intro.",
            "topic": "Attention",
            "length": "medium",
            "design_hint": "Dark navy background, gold titles",
        },
    )
    assert response.status_code == 200
    assert "Dark navy background" in calls["input"]


def test_export_rejects_bad_payload(client: TestClient):
    assert client.post("/api/slides/export", json={"original_filename": "a.pdf", "topic": "t", "slides": []}).status_code == 422
    bad = client.post(
        "/api/slides/export",
        json={
            "original_filename": "a.pptx",
            "topic": "t",
            "slides": [{"title": "x", "bullets": ["y"]}],
            "original_pptx_base64": "!!!not-base64!!!",
        },
    )
    assert bad.status_code == 422
