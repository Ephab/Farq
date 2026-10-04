"""Export a team document as Markdown, Word (.docx) or print-ready HTML (the browser saves it as PDF).

Every format gets the same cover (title, team, course, members, date) and table of contents.
Section content is the team's Markdown, parsed into a small block model so each renderer only
has to know headings, paragraphs and lists. HTML is escaped before any inline formatting is added.
"""
from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Literal
from urllib.parse import quote

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Query
from fastapi.responses import Response
from sqlalchemy import select

from ..identity import CurrentUser, User
from .common import Db, require
from .docs import _document_and_team, _ordered
from .models import Assignment, Course, DocSection, TeamMember
from .policy import authorize

router = APIRouter()

Style = Literal["ieee", "modern"]
FONTS = {"ieee": "Times New Roman", "modern": "Calibri"}
ACCENT = RGBColor(0x2F, 0x5B, 0xD3)
ARABIC = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


@dataclass
class Block:
    kind: Literal["heading", "paragraph", "bullet", "number"]
    text: str
    level: int = 0


@dataclass
class Cover:
    title: str
    team: str
    course: str
    members: list[str]
    issued: str


def parse_markdown(source: str) -> list[Block]:
    """Headings, bullet and numbered lists, and paragraphs (consecutive lines are joined)."""
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph)))
            paragraph.clear()

    for raw in source.splitlines():
        line = raw.strip()
        if not line or line.startswith("```"):
            flush()
            continue
        if heading := re.match(r"(#{1,6})\s+(.*)", line):
            flush()
            blocks.append(Block("heading", heading.group(2), len(heading.group(1))))
        elif bullet := re.match(r"[-*+]\s+(.*)", line):
            flush()
            blocks.append(Block("bullet", bullet.group(1)))
        elif number := re.match(r"\d+[.)]\s+(.*)", line):
            flush()
            blocks.append(Block("number", number.group(1)))
        else:
            paragraph.append(line)
    flush()
    return blocks


def depth(key: str) -> int:
    """Section 1 is a top-level heading, 1.1 one level down, 1.1.1 and deeper share the third."""
    return min(key.count(".") + 1, 3)


def _cover(db: Db, document, team) -> Cover:
    assignment = db.get(Assignment, team.assignment_id) if team.assignment_id else None
    course = db.get(Course, assignment.course_id) if assignment else None
    member_ids = db.scalars(select(TeamMember.user_id).where(TeamMember.team_id == team.id).order_by(TeamMember.joined_at)).all()
    names = []
    for user_id in member_ids:
        user = db.get(User, user_id)
        name = user.display_name if user else "A classmate"
        names.append(f"{name} (lead)" if user_id == team.lead_user_id else name)
    course_line = f"{course.code} — {course.title}" if course else (assignment.title if assignment else "Independent project")
    return Cover(document.title, team.name, course_line, names, date.today().strftime("%d %B %Y"))


# ---------- Markdown ----------

def to_markdown(cover: Cover, sections: list[DocSection]) -> str:
    lines = [f"# {cover.title}", "", f"**Team:** {cover.team}  ", f"**Course:** {cover.course}  ",
             f"**Members:** {', '.join(cover.members)}  ", f"**Date:** {cover.issued}", "", "## Contents", ""]
    lines += [f"{'  ' * (depth(s.key) - 1)}- {s.key} {s.title}" for s in sections]
    for section in sections:
        lines += ["", f"{'#' * (depth(section.key) + 1)} {section.key} {section.title}"]
        if section.content_md.strip():
            lines += ["", section.content_md.strip()]
    return "\n".join(lines) + "\n"


# ---------- HTML (print to PDF) ----------

def _inline_html(text: str) -> str:
    parts = []
    for piece in INLINE.split(text):
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            parts.append(f"<strong>{html.escape(piece[2:-2])}</strong>")
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            parts.append(f"<code>{html.escape(piece[1:-1])}</code>")
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            parts.append(f"<em>{html.escape(piece[1:-1])}</em>")
        else:
            parts.append(html.escape(piece))
    return "".join(parts)


def _blocks_html(blocks: list[Block], base: int) -> str:
    out: list[str] = []
    open_list: str | None = None
    for block in blocks:
        tag = {"bullet": "ul", "number": "ol"}.get(block.kind)
        if open_list and tag != open_list:
            out.append(f"</{open_list}>")
            open_list = None
        if tag and not open_list:
            out.append(f"<{tag}>")
            open_list = tag
        if tag:
            out.append(f'<li dir="auto">{_inline_html(block.text)}</li>')
        elif block.kind == "heading":
            level = min(base + block.level, 6)
            out.append(f'<h{level} dir="auto">{_inline_html(block.text)}</h{level}>')
        else:
            out.append(f'<p dir="auto">{_inline_html(block.text)}</p>')
    if open_list:
        out.append(f"</{open_list}>")
    return "\n".join(out)


PRINT_CSS = """
@page { size: A4; margin: 22mm 20mm; }
body { margin: 0 auto; max-width: 170mm; color: #111; line-height: 1.5; }
body.ieee { font: 11pt/1.45 "Times New Roman", Times, serif; }
body.ieee h1, body.ieee h2, body.ieee h3, body.ieee h4 { font-weight: bold; }
body.modern { font: 10.5pt/1.6 Calibri, "Segoe UI", Arial, sans-serif; }
body.modern h1, body.modern h2 { color: #2f5bd3; }
.cover { min-height: 230mm; display: flex; flex-direction: column; justify-content: center; page-break-after: always; }
.cover h1 { font-size: 26pt; margin: 0 0 18pt; }
body.ieee .cover { text-align: center; }
body.modern .cover { border-left: 6pt solid #2f5bd3; padding-left: 16pt; }
.cover dl { display: grid; grid-template-columns: max-content 1fr; gap: 4pt 12pt; margin: 0; }
body.ieee .cover dl { justify-content: center; text-align: left; }
.cover dt { font-weight: bold; }
.cover dd { margin: 0; }
.toc { page-break-after: always; }
.toc ol { list-style: none; padding: 0; }
.toc li { margin: 3pt 0; }
.toc .d2 { padding-left: 16pt; } .toc .d3 { padding-left: 32pt; }
section h2, section h3, section h4 { page-break-after: avoid; }
code { font-family: Consolas, monospace; font-size: 0.92em; }
@media screen { body { padding: 24px; } }
"""


def to_html(cover: Cover, sections: list[DocSection], style: Style) -> str:
    esc = html.escape
    members = ", ".join(esc(name) for name in cover.members)
    toc = "\n".join(f'<li class="d{depth(s.key)}" dir="auto">{esc(s.key)} {esc(s.title)}</li>' for s in sections)
    body = []
    for section in sections:
        level = depth(section.key) + 1
        body.append(f'<section><h{level} dir="auto">{esc(section.key)} {esc(section.title)}</h{level}>'
                    f"{_blocks_html(parse_markdown(section.content_md), level)}</section>")
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{esc(cover.title)} — {esc(cover.team)}</title>
<style>{PRINT_CSS}</style></head>
<body class="{style}">
<div class="cover"><h1 dir="auto">{esc(cover.title)}</h1><dl>
<dt>Team</dt><dd dir="auto">{esc(cover.team)}</dd><dt>Course</dt><dd dir="auto">{esc(cover.course)}</dd>
<dt>Members</dt><dd dir="auto">{members}</dd><dt>Date</dt><dd>{esc(cover.issued)}</dd></dl></div>
<nav class="toc"><h2>Contents</h2><ol>{toc}</ol></nav>
{"".join(body)}
<script>window.addEventListener("load", () => setTimeout(() => window.print(), 250))</script>
</body></html>"""


# ---------- Word ----------

def _set_font(style_or_run, family: str) -> None:
    """Latin and complex-script (Arabic) text both use the chosen family."""
    style_or_run.font.name = family
    element = style_or_run.element
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), family)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme), None)


def _right_to_left(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _add_inline(paragraph, text: str) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)
    if ARABIC.search(text):
        _right_to_left(paragraph)


def to_docx(cover: Cover, sections: list[DocSection], style: Style) -> bytes:
    word = WordDocument()
    family = FONTS[style]
    normal = word.styles["Normal"]
    _set_font(normal, family)
    normal.font.size = Pt(11 if style == "ieee" else 10.5)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        heading = word.styles[name]
        _set_font(heading, family)
        heading.font.color.rgb = RGBColor(0, 0, 0) if style == "ieee" else ACCENT

    title = word.add_paragraph(cover.title, style="Title")
    for label, value in (("Team", cover.team), ("Course", cover.course), ("Members", ", ".join(cover.members)), ("Date", cover.issued)):
        line = word.add_paragraph()
        line.add_run(f"{label}: ").bold = True
        _add_inline(line, value)
    if style == "ieee":
        for paragraph in word.paragraphs:
            paragraph.alignment = 1
    title.paragraph_format.space_before = Pt(160)
    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    word.add_heading("Contents", level=1)
    for section in sections:
        entry = word.add_paragraph()
        entry.paragraph_format.left_indent = Pt(16 * (depth(section.key) - 1))
        _add_inline(entry, f"{section.key} {section.title}")
    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for section in sections:
        level = depth(section.key)
        heading = word.add_heading(f"{section.key} {section.title}", level=level)
        if ARABIC.search(section.title):
            _right_to_left(heading)
        for block in parse_markdown(section.content_md):
            if block.kind == "heading":
                paragraph = word.add_heading(level=min(level + block.level, 4))
            elif block.kind == "bullet":
                paragraph = word.add_paragraph(style="List Bullet")
            elif block.kind == "number":
                paragraph = word.add_paragraph(style="List Number")
            else:
                paragraph = word.add_paragraph()
            _add_inline(paragraph, block.text)

    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


# ---------- Endpoint ----------

MEDIA = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "html": ("text/html; charset=utf-8", "html"),
}


def _filename(team: str, title: str, extension: str) -> str:
    return f"{team} - {title}.{extension}".replace("/", "-").replace("\\", "-").replace('"', "'")


@router.get("/v1/documents/{document_id}/export")
def export_document(
    document_id: str, db: Db, user: CurrentUser,
    format: Literal["md", "docx", "html"] = Query(), style: Style = Query(default="ieee"),
) -> Response:
    document, team = _document_and_team(db, document_id)
    authorize(db, user, team, "view")
    cover = _cover(db, document, team)
    sections = _ordered(db, document.id)
    if format == "md":
        content: bytes = to_markdown(cover, sections).encode("utf-8")
    elif format == "docx":
        content = to_docx(cover, sections, style)
    else:
        content = to_html(cover, sections, style).encode("utf-8")
    media, extension = MEDIA[format]
    name = _filename(team.name, document.title, extension)
    ascii_name = name.encode("ascii", "ignore").decode() or f"document.{extension}"
    disposition = "inline" if format == "html" else "attachment"
    headers = {"Content-Disposition": f"{disposition}; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}"}
    return Response(content, media_type=media, headers=headers)
